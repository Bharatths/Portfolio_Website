from pathlib import Path
import re

from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parents[1]
MD_FILE = ROOT / "Academic_Project_Report.md"
DOCX_FILE = ROOT / "Academic_Project_Report.docx"
SCREENSHOT_DIR = ROOT / "screenshots"


def clean_inline_markdown(text: str) -> str:
    text = re.sub(r"`([^`]*)`", r"\1", text)
    text = re.sub(r"\*\*([^*]+)\*\*", r"\1", text)
    text = re.sub(r"\*([^*]+)\*", r"\1", text)
    return text.strip()


def add_markdown_to_doc(doc: Document, lines):
    in_code_block = False
    for raw in lines:
        line = raw.rstrip("\n")

        if line.strip().startswith("```"):
            in_code_block = not in_code_block
            continue

        if in_code_block:
            p = doc.add_paragraph(clean_inline_markdown(line))
            p.style = "No Spacing"
            continue

        stripped = line.strip()

        if not stripped:
            doc.add_paragraph("")
            continue

        if stripped == "---":
            doc.add_page_break()
            continue

        if stripped.startswith("### "):
            doc.add_heading(clean_inline_markdown(stripped[4:]), level=3)
            continue

        if stripped.startswith("## "):
            doc.add_heading(clean_inline_markdown(stripped[3:]), level=2)
            continue

        if stripped.startswith("# "):
            h = doc.add_heading(clean_inline_markdown(stripped[2:]), level=1)
            h.alignment = WD_ALIGN_PARAGRAPH.CENTER
            continue

        if stripped.startswith("- "):
            doc.add_paragraph(clean_inline_markdown(stripped[2:]), style="List Bullet")
            continue

        if re.match(r"^\d+\.\s+", stripped):
            content = re.sub(r"^\d+\.\s+", "", stripped)
            doc.add_paragraph(clean_inline_markdown(content), style="List Number")
            continue

        doc.add_paragraph(clean_inline_markdown(stripped))


def add_screenshots_section(doc: Document):
    doc.add_page_break()
    doc.add_heading("Output Screens (Attached)", level=1)

    expected = [
        "01-home.png",
        "02-about.png",
        "03-skills.png",
        "04-projects-certifications.png",
        "05-contact-form.png",
        "06-mobile-navbar.png",
    ]

    if not SCREENSHOT_DIR.exists():
        doc.add_paragraph(
            "Screenshot folder not found. Create 'screenshots' folder in project root and add the required images."
        )
        for name in expected:
            doc.add_paragraph(name, style="List Bullet")
        return

    available = []
    for name in expected:
        file_path = SCREENSHOT_DIR / name
        if file_path.exists():
            available.append(file_path)

    if not available:
        doc.add_paragraph(
            "No expected screenshots found. Add images with the names below and run this script again:"
        )
        for name in expected:
            doc.add_paragraph(name, style="List Bullet")
        return

    for image in available:
        caption = image.stem.replace("-", " ").title()
        doc.add_heading(caption, level=2)
        try:
            doc.add_picture(str(image), width=Inches(6.3))
        except Exception:
            doc.add_paragraph(f"Could not insert image: {image.name}")


if __name__ == "__main__":
    if not MD_FILE.exists():
        raise FileNotFoundError(f"Report file not found: {MD_FILE}")

    doc = Document()
    lines = MD_FILE.read_text(encoding="utf-8").splitlines(True)
    add_markdown_to_doc(doc, lines)
    add_screenshots_section(doc)
    doc.save(DOCX_FILE)

    print(f"Generated: {DOCX_FILE}")
